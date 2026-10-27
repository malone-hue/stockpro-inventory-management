import os

import mysql.connector
from flask import (
    Flask,
    jsonify,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session
)
from dotenv import load_dotenv
from auth import auth, login_required, admin_required
from flask_wtf.csrf import CSRFProtect
from io import BytesIO
from openpyxl import Workbook
from flask import send_file
from pathlib import Path

load_dotenv()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("FLASK_SECRET_KEY")

csrf = CSRFProtect(app)

# Register authentication routes
app.register_blueprint(auth)


def get_db_connection():
    ca_path = Path(
        os.getenv("DB_SSL_CA", "certs/tidb-ca.pem")
    ).resolve()

    if not ca_path.is_file():
        raise FileNotFoundError(
            f"TiDB CA certificate not found: {ca_path}"
        )

    return mysql.connector.connect(
        host=os.getenv("DB_HOST"),
        port=int(os.getenv("DB_PORT", "4000")),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD"),
        database=os.getenv("DB_NAME"),
        ssl_ca=str(ca_path),
        ssl_verify_cert=True,
        ssl_verify_identity=True,
        connection_timeout=15
    )




@app.route("/")
@login_required
def home():
    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)
        # Get the 10 most recent stock movements

        cursor.execute("""
            SELECT
                sm.id,
                p.name AS product_name,
                sm.movement_type,
                sm.quantity,
                sm.reason,
                sm.created_at
            FROM stock_movements sm
            JOIN products p ON sm.product_id = p.id
            ORDER BY sm.created_at DESC, sm.id DESC
            LIMIT 10
        """)

        recent_movements = cursor.fetchall()

        # Get search and category filters.
        search_query = request.args.get("search", "").strip()
        selected_category = request.args.get("category", "").strip()

        # Get all available categories.
        cursor.execute(
            "SELECT DISTINCT category FROM products ORDER BY category ASC"
        )
        categories = [row["category"] for row in cursor.fetchall()]

        # Get filtered products for the inventory table.
        query = "SELECT * FROM products WHERE 1=1"
        parameters = []

        if search_query:
            query += " AND name LIKE %s"
            parameters.append(f"%{search_query}%")

        if selected_category:
            query += " AND category = %s"
            parameters.append(selected_category)

        query += " ORDER BY id DESC"

        cursor.execute(query, tuple(parameters))
        products = cursor.fetchall()

    
        # Calculate today's sales revenue and gross profit.
        cursor.execute("""
            SELECT
                COALESCE(SUM(total_amount), 0) AS todays_sales,
                COALESCE(SUM(gross_profit), 0) AS todays_gross_profit,
                COUNT(*) AS todays_transactions
            FROM sales
            WHERE created_at >= CURDATE()
              AND created_at < CURDATE() + INTERVAL 1 DAY
        """)

        sales_analytics = cursor.fetchone()

        todays_sales = float(sales_analytics["todays_sales"])
        todays_gross_profit = float(sales_analytics["todays_gross_profit"])
        todays_transactions = int(sales_analytics["todays_transactions"])



        # Get ALL products for dashboard analytics and alerts.
        cursor.execute("SELECT * FROM products ORDER BY id DESC")
        all_products = cursor.fetchall()

        # Calculate low-stock and out-of-stock products across the inventory.
        low_stock_products = [
            product for product in all_products
            if product["quantity"] <= product["reorder_level"]
        ]

        out_of_stock_count = sum(
            1 for product in all_products
            if product["quantity"] <= 0
        )

        low_stock_count = len(low_stock_products)
        total_products = len(all_products)

        total_quantity = sum(
            int(product["quantity"]) for product in all_products
        )

        total_cost_value = sum(
            float(product["cost_price"]) * int(product["quantity"])
            for product in all_products
        )

        # Potential sales value and gross profit at current selling prices.
        potential_sales_value = sum(
            float(product["selling_price"]) * int(product["quantity"])
            for product in all_products
        )

        potential_gross_profit = sum(
            (float(product["selling_price"]) - float(product["cost_price"]))
            * int(product["quantity"])
            for product in all_products
        )

        # Prepare chart data using ALL products, not filtered products.
        category_totals = {}

        for product in all_products:
            category = product["category"] or "Uncategorized"

            category_totals[category] = (
                category_totals.get(category, 0)
                + int(product["quantity"])
            )

        chart_products = all_products[:10]

        return render_template(
            "index.html",
            products=products,
            recent_movements=recent_movements,
            all_products=all_products,
            categories=categories,
            search_query=search_query,
            selected_category=selected_category,
            low_stock_count=low_stock_count,
            low_stock_products=low_stock_products,
            out_of_stock_count=out_of_stock_count,
            total_products=total_products,
            total_quantity=total_quantity,
            total_cost_value=total_cost_value,
            potential_sales_value=potential_sales_value,
            potential_gross_profit=potential_gross_profit,
            category_totals=category_totals,

            todays_sales=todays_sales,
            todays_gross_profit=todays_gross_profit,
            todays_transactions=todays_transactions,


            chart_products=chart_products
        )

    except mysql.connector.Error:
        app.logger.exception("Could not load products")
        return "Database error. Check the VS Code terminal.", 500

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()


@app.route("/products/add", methods=["GET", "POST"])
@admin_required
def add_product():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        category = request.form.get("category", "").strip()

        try:
            cost_price = float(request.form.get("cost_price", 0))
            selling_price = float(request.form.get("selling_price", 0))
            quantity = int(request.form.get("quantity", 0))
            reorder_level = int(request.form.get("reorder_level", 5))

            if not name or not category:
                flash("Product name and category are required.", "error")
                return redirect(url_for("add_product"))

            if (
                cost_price < 0 or selling_price < 0
                or quantity < 0 or reorder_level < 0
            ):
                flash(
                    "Prices, quantity, and reorder level cannot be negative.",
                    "error"
                )
                return redirect(url_for("add_product"))

        except (ValueError, TypeError):
            flash(
                "Please enter valid numbers for prices and quantities.",
                "error"
            )
            return redirect(url_for("add_product"))

        connection = None
        cursor = None

        try:
            connection = get_db_connection()
            cursor = connection.cursor()

            cursor.execute(
                """
                INSERT INTO products
                (name, category, cost_price, selling_price,
                 quantity, reorder_level)
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (
                    name, category, cost_price,
                    selling_price, quantity, reorder_level
                )
            )

            connection.commit()
            flash("Product added successfully!", "success")
            return redirect(url_for("home"))

        except mysql.connector.Error:
            if connection is not None:
                connection.rollback()

            app.logger.exception("Could not add product")
            flash(
                "Could not save the product. Check the server terminal.",
                "error"
            )

        finally:
            if cursor is not None:
                cursor.close()

            if connection is not None and connection.is_connected():
                connection.close()

    return render_template("add_product.html")


@app.route("/products/edit/<int:product_id>", methods=["GET", "POST"])
@admin_required
def edit_product(product_id):
    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        if request.method == "POST":
            name = request.form.get("name", "").strip()
            category = request.form.get("category", "").strip()

            try:
                cost_price = float(request.form.get("cost_price", ""))
                selling_price = float(request.form.get("selling_price", ""))
                quantity = int(request.form.get("quantity", ""))
                reorder_level = int(request.form.get("reorder_level", ""))

            except ValueError:
                flash(
                    "Enter valid prices and whole numbers for quantities.",
                    "error"
                )
                return redirect(
                    url_for("edit_product", product_id=product_id)
                )

            if not name or not category:
                flash(
                    "Product name and category are required.",
                    "error"
                )
                return redirect(
                    url_for("edit_product", product_id=product_id)
                )

            if (
                cost_price < 0 or selling_price < 0
                or quantity < 0 or reorder_level < 0
            ):
                flash("Values cannot be negative.", "error")
                return redirect(
                    url_for("edit_product", product_id=product_id)
                )

            cursor.execute(
                """
                UPDATE products
                SET name = %s,
                    category = %s,
                    cost_price = %s,
                    selling_price = %s,
                    quantity = %s,
                    reorder_level = %s
                WHERE id = %s
                """,
                (
                    name, category, cost_price,
                    selling_price, quantity,
                    reorder_level, product_id
                )
            )

            connection.commit()

            if cursor.rowcount == 0:
                cursor.execute(
                    "SELECT id FROM products WHERE id = %s",
                    (product_id,)
                )

                if cursor.fetchone() is None:
                    flash("Product not found.", "error")
                    return redirect(url_for("home"))

            flash("Product updated successfully!", "success")
            return redirect(url_for("home"))

        cursor.execute(
            "SELECT * FROM products WHERE id = %s",
            (product_id,)
        )
        product = cursor.fetchone()

        if product is None:
            flash("Product not found.", "error")
            return redirect(url_for("home"))

        return render_template("edit_product.html", product=product)

    except mysql.connector.Error:
        if connection is not None:
            connection.rollback()

        app.logger.exception("Could not edit product")
        return "A database error occurred. Check the VS Code terminal.", 500

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()


@app.route("/products/delete/<int:product_id>", methods=["POST"])
@admin_required
def delete_product(product_id):
    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor()

        cursor.execute(
            "DELETE FROM products WHERE id = %s",
            (product_id,)
        )
        connection.commit()

        if cursor.rowcount == 0:
            flash("Product not found.", "error")
        else:
            flash("Product deleted successfully.", "success")

        return redirect(url_for("home"))

    except mysql.connector.Error:
        if connection is not None:
            connection.rollback()

        app.logger.exception("Could not delete product")
        return "A database error occurred. Check the VS Code terminal.", 500

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()


@app.route("/api/products")
@login_required
def get_products():
    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        cursor.execute("SELECT * FROM products ORDER BY id DESC")
        products = cursor.fetchall()

        return jsonify(products)

    except mysql.connector.Error:
        app.logger.exception(
            "Database error while fetching products"
        )
        return jsonify({"error": "Could not retrieve products"}), 500

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()

@app.route("/stock-movement", methods=["GET", "POST"])
@login_required
def stock_movement():
    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        if request.method == "POST":
            product_id = request.form.get("product_id", type=int)
            movement_type = request.form.get("movement_type", "").upper()
            quantity = request.form.get("quantity", type=int)
            reason = request.form.get("reason", "").strip()

            if not product_id or quantity is None or quantity <= 0:
                flash("Please select a product and enter a valid quantity.", "error")
                return redirect(url_for("stock_movement"))

            if movement_type not in ("IN", "OUT"):
                flash("Invalid stock movement type.", "error")
                return redirect(url_for("stock_movement"))

            # Start a database transaction.
            connection.start_transaction()

            # Lock the product row while updating its stock.
            cursor.execute(
                "SELECT id, name, quantity FROM products WHERE id = %s FOR UPDATE",
                (product_id,)
            )
            product = cursor.fetchone()

            if product is None:
                connection.rollback()
                flash("Product not found.", "error")
                return redirect(url_for("stock_movement"))

            current_quantity = product["quantity"]

            if movement_type == "OUT" and quantity > current_quantity:
                connection.rollback()
                flash(
                    f"Insufficient stock. Only {current_quantity} units are available.",
                    "error"
                )
                return redirect(url_for("stock_movement"))

            if movement_type == "IN":
                new_quantity = current_quantity + quantity
            else:
                new_quantity = current_quantity - quantity

            # Update the product quantity.
            cursor.execute(
                "UPDATE products SET quantity = %s WHERE id = %s",
                (new_quantity, product_id)
            )

            # Record the movement in the history table.
            cursor.execute(
                """
                INSERT INTO stock_movements
                    (product_id, movement_type, quantity, reason)
                VALUES (%s, %s, %s, %s)
                """,
                (product_id, movement_type, quantity, reason or None)
            )

            # Save both changes together.
            connection.commit()

            action = "added to" if movement_type == "IN" else "removed from"
            flash(
                f"{quantity} units {action} {product['name']} successfully.",
                "success"
            )

            return redirect(url_for("stock_movement"))

        # Load products for the form.
        cursor.execute(
            "SELECT id, name, quantity FROM products ORDER BY name"
        )
        products = cursor.fetchall()

        # Load recent stock movement history.
        cursor.execute(
            """
            SELECT
                sm.id,
                p.name AS product_name,
                sm.movement_type,
                sm.quantity,
                sm.reason,
                sm.created_at
            FROM stock_movements sm
            JOIN products p ON sm.product_id = p.id
            ORDER BY sm.created_at DESC, sm.id DESC
            LIMIT 100
            """
        )
        movements = cursor.fetchall()

        return render_template(
            "stock_movement.html",
            products=products,
            movements=movements
        )

    except Exception:
        if connection is not None:
            connection.rollback()

        app.logger.exception("Stock movement operation failed")
        flash("Something went wrong while processing stock. Please try again.", "error")
        return redirect(url_for("stock_movement"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None:
            connection.close()

@app.route("/sales/new", methods=["GET", "POST"])
@login_required
def record_sale():
    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        if request.method == "POST":
            product_id = request.form.get("product_id", type=int)
            quantity = request.form.get("quantity", type=int)

            if not product_id or quantity is None or quantity <= 0:
                flash("Select a product and enter a valid quantity.", "error")
                return redirect(url_for("record_sale"))

            # Begin a transaction so all sale-related changes succeed together.
            connection.start_transaction()

            # Lock the selected product to prevent conflicting stock updates.
            cursor.execute(
                """
                SELECT id, name, quantity, cost_price, selling_price
                FROM products
                WHERE id = %s
                FOR UPDATE
                """,
                (product_id,)
            )
            product = cursor.fetchone()

            if product is None:
                connection.rollback()
                flash("Product not found.", "error")
                return redirect(url_for("record_sale"))

            if quantity > product["quantity"]:
                connection.rollback()
                flash(
                    f"Insufficient stock. Only {product['quantity']} "
                    "units are available.",
                    "error"
                )
                return redirect(url_for("record_sale"))

            cost_price = product["cost_price"]
            selling_price = product["selling_price"]

            total_amount = selling_price * quantity
            gross_profit = (selling_price - cost_price) * quantity
            new_quantity = product["quantity"] - quantity

            # Deduct the sold quantity from inventory.
            cursor.execute(
                "UPDATE products SET quantity = %s WHERE id = %s",
                (new_quantity, product_id)
            )

            # Record the sale using the prices at the time of sale.
            cursor.execute(
                """
                INSERT INTO sales (
                    product_id,
                    quantity,
                    cost_price,
                    selling_price,
                    total_amount,
                    gross_profit,
                    sold_by
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    product_id,
                    quantity,
                    cost_price,
                    selling_price,
                    total_amount,
                    gross_profit,
                    session.get("user_id")
                )
            )

            sale_id = cursor.lastrowid

            # Record the corresponding stock-out movement.
            cursor.execute(
                """
                INSERT INTO stock_movements (
                    product_id,
                    movement_type,
                    quantity,
                    reason
                )
                VALUES (%s, 'OUT', %s, %s)
                """,
                (product_id, quantity, f"Sale #{sale_id}")
            )

            # Commit all three database changes together.
            connection.commit()

            flash(
                f"Sale #{sale_id} recorded successfully! "
                f"Total: {total_amount:.2f}; "
                f"Gross profit: {gross_profit:.2f}.",
                "success"
            )

            return redirect(url_for("record_sale"))

        # Load products available for selection.
        cursor.execute(
            """
            SELECT id, name, quantity, selling_price
            FROM products
            WHERE quantity > 0
            ORDER BY name
            """
        )
        products = cursor.fetchall()

        # Load recent sales.
        cursor.execute(
            """
            SELECT
                s.id,
                p.name AS product_name,
                s.quantity,
                s.selling_price,
                s.total_amount,
                s.gross_profit,
                s.created_at
            FROM sales s
            JOIN products p ON s.product_id = p.id
            ORDER BY s.created_at DESC, s.id DESC
            LIMIT 50
            """
        )
        sales = cursor.fetchall()

        return render_template(
            "sales.html",
            products=products,
            sales=sales
        )

    except Exception:
        if connection is not None:
            connection.rollback()

        app.logger.exception("Sale operation failed")
        flash("Unable to process the sale. Please try again.", "error")
        return redirect(url_for("record_sale"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None:
            connection.close()




@app.route("/stock-history")
@login_required
def stock_history():
    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        cursor.execute(
            """
            SELECT
                sm.id,
                p.name AS product_name,
                sm.movement_type,
                sm.quantity,
                sm.reason,
                sm.created_at
            FROM stock_movements sm
            JOIN products p ON sm.product_id = p.id
            ORDER BY sm.created_at DESC, sm.id DESC
            """
        )

        movements = cursor.fetchall()

        return render_template(
            "stock_history.html",
            movements=movements
        )

    except Exception:
        app.logger.exception("Could not load stock history")
        flash("Could not load stock history.", "error")
        return redirect(url_for("home"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None:
            connection.close()


@app.route("/reports")
@admin_required
def reports():
    connection = None
    cursor = None

    # Read date filters from the URL.
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()

    try:
        from datetime import date

        # Validate the supplied dates.
        if start_date:
            date.fromisoformat(start_date)

        if end_date:
            date.fromisoformat(end_date)

        if start_date and end_date and start_date > end_date:
            flash("Start date cannot be after end date.", "error")
            return redirect(url_for("reports"))

        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)
        
        # Sales and profit summary for the selected date range.
        sales_query = """
            SELECT
                COALESCE(SUM(s.total_amount), 0) AS total_revenue,
                COALESCE(
                    SUM(s.cost_price * s.quantity), 0
                ) AS cost_of_goods_sold,
                COALESCE(SUM(s.gross_profit), 0) AS gross_profit,
                COUNT(*) AS total_transactions,
                COALESCE(SUM(s.quantity), 0) AS units_sold
            FROM sales s
            WHERE 1 = 1
        """

        sales_params = []

        if start_date:
            sales_query += " AND s.created_at >= %s"
            sales_params.append(start_date)

        if end_date:
            sales_query += """
                AND s.created_at < DATE_ADD(%s, INTERVAL 1 DAY)
            """
            sales_params.append(end_date)

        cursor.execute(sales_query, tuple(sales_params))
        sales_summary = cursor.fetchone()



        # Overall inventory summary.
        cursor.execute("""
            SELECT
                COUNT(*) AS total_products,
                COALESCE(SUM(quantity), 0) AS total_quantity,
                COALESCE(SUM(quantity * cost_price), 0)
                    AS inventory_cost_value,
                COALESCE(SUM(quantity * selling_price), 0)
                    AS potential_sales_value,
                COALESCE(
                    SUM(quantity * (selling_price - cost_price)), 0
                ) AS potential_gross_profit
            FROM products
        """)

        summary = cursor.fetchone()

        # Product inventory report.
        cursor.execute("""
            SELECT
                id,
                name,
                category,
                cost_price,
                selling_price,
                quantity,
                reorder_level,
                quantity * cost_price AS stock_cost_value,
                quantity * selling_price AS stock_sales_value
            FROM products
            ORDER BY
                CASE WHEN quantity <= reorder_level THEN 0 ELSE 1 END,
                name ASC
        """)

        products = cursor.fetchall()

        low_stock_products = [
            product for product in products
            if product["quantity"] <= product["reorder_level"]
        ]

        # Build the stock movement date filter safely.
        movement_query = """
            SELECT
                sm.id,
                p.name AS product_name,
                sm.movement_type,
                sm.quantity,
                sm.reason,
                sm.created_at
            FROM stock_movements sm
            JOIN products p ON p.id = sm.product_id
            WHERE 1 = 1
        """

        movement_params = []

        if start_date:
            movement_query += " AND sm.created_at >= %s"
            movement_params.append(start_date)

        if end_date:
            movement_query += " AND sm.created_at < DATE_ADD(%s, INTERVAL 1 DAY)"
            movement_params.append(end_date)

        movement_query += " ORDER BY sm.created_at DESC, sm.id DESC"

        cursor.execute(movement_query, tuple(movement_params))
        movements = cursor.fetchall()

        # Calculate movement totals for the selected period.
        stock_in_quantity = sum(
            movement["quantity"]
            for movement in movements
            if movement["movement_type"] == "IN"
        )

        stock_out_quantity = sum(
            movement["quantity"]
            for movement in movements
            if movement["movement_type"] == "OUT"
        )

        return render_template(
            "reports.html",
            summary=summary,
            sales_summary=sales_summary,
            products=products,
            low_stock_products=low_stock_products,
            movements=movements,
            stock_in_quantity=stock_in_quantity,
            stock_out_quantity=stock_out_quantity,
            start_date=start_date,
            
            end_date=end_date
        )

    except ValueError:
        flash("Please enter valid dates.", "error")
        return redirect(url_for("reports"))

    except Exception:
        app.logger.exception("Could not generate inventory reports")
        flash("Unable to generate reports. Please try again.", "error")
        return redirect(url_for("home"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()

    
@app.route("/reports/export/pdf")
@admin_required
def export_reports_pdf():
    connection = None
    cursor = None

    try:
        from datetime import date, datetime
        from io import BytesIO

        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
            PageBreak,
        )

        # Read and validate the date filters.
        start_date = request.args.get("start_date", "").strip()
        end_date = request.args.get("end_date", "").strip()

        if start_date:
            date.fromisoformat(start_date)

        if end_date:
            date.fromisoformat(end_date)

        if start_date and end_date and start_date > end_date:
            flash("Start date cannot be after end date.", "error")
            return redirect(url_for("reports"))

        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        # Reuse the same date filtering logic as the Excel export.
        sales_query = """
            SELECT
                COALESCE(SUM(s.total_amount), 0) AS total_revenue,
                COALESCE(
                    SUM(s.cost_price * s.quantity), 0
                ) AS cost_of_goods_sold,
                COALESCE(SUM(s.gross_profit), 0) AS gross_profit,
                COUNT(*) AS total_transactions,
                COALESCE(SUM(s.quantity), 0) AS units_sold
            FROM sales s
            WHERE 1 = 1
        """

        sales_params = []

        if start_date:
            sales_query += " AND s.created_at >= %s"
            sales_params.append(start_date)

        if end_date:
            sales_query += """
                AND s.created_at < DATE_ADD(%s, INTERVAL 1 DAY)
            """
            sales_params.append(end_date)

        cursor.execute(sales_query, tuple(sales_params))
        sales_summary = cursor.fetchone()

        # Get current inventory.
        cursor.execute("""
            SELECT
                id,
                name,
                category,
                cost_price,
                selling_price,
                quantity,
                reorder_level
            FROM products
            ORDER BY name ASC
        """)
        products = cursor.fetchall()

        # Get stock movements for the selected period.
        movement_query = """
            SELECT
                sm.id,
                sm.created_at,
                p.name AS product_name,
                sm.movement_type,
                sm.quantity,
                sm.reason
            FROM stock_movements sm
            JOIN products p ON p.id = sm.product_id
            WHERE 1 = 1
        """

        movement_params = []

        if start_date:
            movement_query += " AND sm.created_at >= %s"
            movement_params.append(start_date)

        if end_date:
            movement_query += """
                AND sm.created_at < DATE_ADD(%s, INTERVAL 1 DAY)
            """
            movement_params.append(end_date)

        movement_query += " ORDER BY sm.created_at DESC, sm.id DESC"

        cursor.execute(movement_query, tuple(movement_params))
        movements = cursor.fetchall()

        # Prepare the PDF.
        output = BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
            title="StockPro Business Report",
            author="StockPro",
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "StockProTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#1D4ED8"),
            fontSize=20,
            spaceAfter=8,
        )

        section_style = ParagraphStyle(
            "StockProSection",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#1D4ED8"),
            spaceBefore=12,
            spaceAfter=8,
        )

        small_style = ParagraphStyle(
            "StockProSmall",
            parent=styles["BodyText"],
            fontSize=8,
            leading=10,
        )

        story = []

        story.append(Paragraph("StockPro Business Report", title_style))

        period = (
            f"{start_date or 'All dates'} to {end_date or 'All dates'}"
        )

        story.append(
            Paragraph(f"<b>Reporting period:</b> {period}", styles["Normal"])
        )
        story.append(
            Paragraph(
                f"<b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}",
                styles["Normal"],
            )
        )
        story.append(Spacer(1, 8 * mm))

        # Sales summary section.
        story.append(Paragraph("Sales Summary", section_style))

        sales_data = [
            ["Metric", "Value"],
            [
                "Total Sales Revenue",
                f"KSh {float(sales_summary['total_revenue']):,.2f}",
            ],
            [
                "Cost of Goods Sold",
                f"KSh {float(sales_summary['cost_of_goods_sold']):,.2f}",
            ],
            [
                "Gross Profit",
                f"KSh {float(sales_summary['gross_profit']):,.2f}",
            ],
            [
                "Total Transactions",
                str(sales_summary["total_transactions"]),
            ],
            ["Units Sold", str(sales_summary["units_sold"])],
        ]

        sales_table = Table(sales_data, colWidths=[75 * mm, 55 * mm])
        sales_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ("PADDING", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))

        story.append(sales_table)
        story.append(Spacer(1, 6 * mm))

        # Inventory section.
        story.append(Paragraph("Current Inventory", section_style))

        inventory_data = [[
            "ID",
            "Product",
            "Category",
            "Cost Price",
            "Selling Price",
            "Quantity",
            "Reorder Level",
            "Stock Cost Value",
            "Potential Sales Value",
        ]]

        total_stock_cost = 0
        total_potential_sales = 0

        for product in products:
            quantity = int(product["quantity"])
            cost_price = float(product["cost_price"])
            selling_price = float(product["selling_price"])

            stock_cost = quantity * cost_price
            stock_sales = quantity * selling_price

            total_stock_cost += stock_cost
            total_potential_sales += stock_sales

            inventory_data.append([
                str(product["id"]),
                Paragraph(str(product["name"]), small_style),
                Paragraph(str(product["category"] or ""), small_style),
                f"{cost_price:,.2f}",
                f"{selling_price:,.2f}",
                str(quantity),
                str(product["reorder_level"]),
                f"{stock_cost:,.2f}",
                f"{stock_sales:,.2f}",
            ])

        inventory_data.append([
            "",
            "TOTAL",
            "",
            "",
            "",
            str(sum(int(p["quantity"]) for p in products)),
            "",
            f"{total_stock_cost:,.2f}",
            f"{total_potential_sales:,.2f}",
        ])

        inventory_table = Table(
            inventory_data,
            repeatRows=1,
            colWidths=[
                12 * mm, 36 * mm, 25 * mm, 22 * mm, 24 * mm,
                17 * mm, 20 * mm, 27 * mm, 30 * mm,
            ],
        )

        inventory_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("PADDING", (0, 0), (-1, -1), 4),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DBEAFE")),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ]))

        story.append(inventory_table)
        story.append(PageBreak())

        # Stock movement section.
        story.append(Paragraph("Stock Movements", section_style))

        movement_data = [[
            "Movement ID",
            "Date",
            "Product",
            "Type",
            "Quantity",
            "Reason",
        ]]

        for movement in movements:
            created_at = movement["created_at"]

            movement_data.append([
                str(movement["id"]),
                str(created_at),
                Paragraph(str(movement["product_name"]), small_style),
                str(movement["movement_type"]),
                str(movement["quantity"]),
                Paragraph(
                    str(movement["reason"] or "Not specified"),
                    small_style,
                ),
            ])

        movement_table = Table(
            movement_data,
            repeatRows=1,
            colWidths=[
                22 * mm, 40 * mm, 48 * mm,
                28 * mm, 25 * mm, 65 * mm,
            ],
        )

        movement_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("PADDING", (0, 0), (-1, -1), 5),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))

        story.append(movement_table)

        # Add page numbers to each page.
        def add_page_number(canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#6B7280"))
            canvas.drawRightString(
                landscape(A4)[0] - 12 * mm,
                8 * mm,
                f"StockPro | Page {doc.page}",
            )
            canvas.restoreState()

        document.build(
            story,
            onFirstPage=add_page_number,
            onLaterPages=add_page_number,
        )

        output.seek(0)

        filename = "StockPro_Report"

        if start_date:
            filename += f"_{start_date}"

        if end_date:
            filename += f"_to_{end_date}"

        filename += ".pdf"

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/pdf",
        )

    except ValueError:
        flash("Please enter valid dates.", "error")
        return redirect(url_for("reports"))

    except Exception:
        app.logger.exception("Could not export PDF report")
        flash("Unable to export the PDF report.", "error")
        return redirect(url_for("reports"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()


            
@app.route("/reports/export/excel")
@admin_required
def export_reports_excel():
    connection = None
    cursor = None

    try:
        from datetime import date
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        # Read and validate the date filters.
        start_date = request.args.get("start_date", "").strip()
        end_date = request.args.get("end_date", "").strip()

        if start_date:
            date.fromisoformat(start_date)

        if end_date:
            date.fromisoformat(end_date)

        if start_date and end_date and start_date > end_date:
            flash("Start date cannot be after end date.", "error")
            return redirect(url_for("reports"))

        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        workbook = Workbook()

        # --------------------------------------------------
        # Worksheet 1: Sales Summary
        # --------------------------------------------------
        sales_sheet = workbook.active
        sales_sheet.title = "Sales Summary"

        sales_query = """
            SELECT
                COALESCE(SUM(s.total_amount), 0) AS total_revenue,
                COALESCE(
                    SUM(s.cost_price * s.quantity), 0
                ) AS cost_of_goods_sold,
                COALESCE(SUM(s.gross_profit), 0) AS gross_profit,
                COUNT(*) AS total_transactions,
                COALESCE(SUM(s.quantity), 0) AS units_sold
            FROM sales s
            WHERE 1 = 1
        """

        sales_params = []

        if start_date:
            sales_query += " AND s.created_at >= %s"
            sales_params.append(start_date)

        if end_date:
            sales_query += """
                AND s.created_at < DATE_ADD(%s, INTERVAL 1 DAY)
            """
            sales_params.append(end_date)

        cursor.execute(sales_query, tuple(sales_params))
        sales_summary = cursor.fetchone()

        sales_sheet.append(["StockPro Sales Report"])
        sales_sheet.append(["Start Date", start_date or "All dates"])
        sales_sheet.append(["End Date", end_date or "All dates"])
        sales_sheet.append([])
        sales_sheet.append(["Metric", "Value"])

        metrics = [
            ("Total Sales Revenue", float(sales_summary["total_revenue"])),
            ("Cost of Goods Sold", float(sales_summary["cost_of_goods_sold"])),
            ("Gross Profit", float(sales_summary["gross_profit"])),
            ("Total Transactions", int(sales_summary["total_transactions"])),
            ("Units Sold", int(sales_summary["units_sold"])),
        ]

        for metric in metrics:
            sales_sheet.append(list(metric))

        # --------------------------------------------------
        # Worksheet 2: Inventory
        # --------------------------------------------------
        inventory_sheet = workbook.create_sheet("Inventory")

        inventory_sheet.append([
            "Product ID",
            "Product Name",
            "Category",
            "Cost Price",
            "Selling Price",
            "Quantity",
            "Reorder Level",
            "Stock Cost Value",
            "Potential Sales Value",
        ])

        cursor.execute("""
            SELECT
                id,
                name,
                category,
                cost_price,
                selling_price,
                quantity,
                reorder_level
            FROM products
            ORDER BY name ASC
        """)

        for product in cursor.fetchall():
            quantity = int(product["quantity"])
            cost_price = float(product["cost_price"])
            selling_price = float(product["selling_price"])

            inventory_sheet.append([
                product["id"],
                product["name"],
                product["category"],
                cost_price,
                selling_price,
                quantity,
                product["reorder_level"],
                cost_price * quantity,
                selling_price * quantity,
            ])

        # --------------------------------------------------
        # Worksheet 3: Stock Movements
        # --------------------------------------------------
        movements_sheet = workbook.create_sheet("Stock Movements")

        movements_sheet.append([
            "Movement ID",
            "Date",
            "Product",
            "Movement Type",
            "Quantity",
            "Reason",
        ])

        movement_query = """
            SELECT
                sm.id,
                sm.created_at,
                p.name AS product_name,
                sm.movement_type,
                sm.quantity,
                sm.reason
            FROM stock_movements sm
            JOIN products p ON p.id = sm.product_id
            WHERE 1 = 1
        """

        movement_params = []

        if start_date:
            movement_query += " AND sm.created_at >= %s"
            movement_params.append(start_date)

        if end_date:
            movement_query += """
                AND sm.created_at < DATE_ADD(%s, INTERVAL 1 DAY)
            """
            movement_params.append(end_date)

        movement_query += " ORDER BY sm.created_at DESC, sm.id DESC"

        cursor.execute(movement_query, tuple(movement_params))

        for movement in cursor.fetchall():
            movements_sheet.append([
                movement["id"],
                movement["created_at"],
                movement["product_name"],
                movement["movement_type"],
                movement["quantity"],
                movement["reason"] or "Not specified",
            ])

        # --------------------------------------------------
        # Format worksheets
        # --------------------------------------------------
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="1D4ED8"
                )
                cell.alignment = Alignment(
                    horizontal="center"
                )

            for column_cells in sheet.columns:
                max_length = 0
                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:
                    value = str(cell.value or "")
                    max_length = max(max_length, len(value))

                sheet.column_dimensions[column_letter].width = min(
                    max_length + 2, 35
                )

        # Highlight the Sales Summary heading.
        sales_sheet["A1"].font = Font(
            bold=True,
            size=16,
            color="1D4ED8"
        )

        # Format monetary values.
        for row in range(2, sales_sheet.max_row + 1):
            label = sales_sheet.cell(row=row, column=1).value

            if label in {
                "Total Sales Revenue",
                "Cost of Goods Sold",
                "Gross Profit",
            }:
                sales_sheet.cell(
                    row=row, column=2
                ).number_format = '"KSh" #,##0.00'

        for row in range(2, inventory_sheet.max_row + 1):
            for column in (4, 5, 8, 9):
                inventory_sheet.cell(
                    row=row, column=column
                ).number_format = '"KSh" #,##0.00'

        # Write the workbook into memory.
        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        filename = "StockPro_Report"

        if start_date:
            filename += f"_{start_date}"

        if end_date:
            filename += f"_to_{end_date}"

        filename += ".xlsx"

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    except ValueError:
        flash("Please enter valid dates.", "error")
        return redirect(url_for("reports"))

    except Exception:
        app.logger.exception("Could not export Excel report")
        flash("Unable to export the Excel report.", "error")
        return redirect(url_for("reports"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()





@app.route("/reports/export")
@admin_required
def export_inventory_report():
    import csv
    import io
    from flask import Response

    connection = None
    cursor = None

    try:
        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        cursor.execute("""
            SELECT
                id,
                name,
                category,
                cost_price,
                selling_price,
                quantity,
                reorder_level,
                quantity * cost_price AS stock_cost_value,
                quantity * selling_price AS stock_sales_value
            FROM products
            ORDER BY name ASC
        """)

        products = cursor.fetchall()

        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow([
            "ID",
            "Product Name",
            "Category",
            "Cost Price",
            "Selling Price",
            "Quantity",
            "Reorder Level",
            "Stock Cost Value",
            "Potential Sales Value",
            "Stock Status"
        ])

        for product in products:
            status = (
                "LOW STOCK"
                if product["quantity"] <= product["reorder_level"]
                else "IN STOCK"
            )

            writer.writerow([
                product["id"],
                product["name"],
                product["category"],
                product["cost_price"],
                product["selling_price"],
                product["quantity"],
                product["reorder_level"],
                product["stock_cost_value"],
                product["stock_sales_value"],
                status
            ])

        csv_content = "\ufeff" + output.getvalue()
        output.close()

        return Response(
            csv_content,
            mimetype="text/csv",
            headers={
                "Content-Disposition":
                    "attachment; filename=inventory_report.csv"
            }
        )

    except Exception:
        app.logger.exception("Could not export inventory report")
        flash("Unable to export the inventory report.", "error")
        return redirect(url_for("reports"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None:
            connection.close()


@app.route("/reports/export-movements")
@admin_required
def export_stock_movement_report():
    import csv
    import io
    from flask import Response
    from datetime import date

    connection = None
    cursor = None

    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()

    try:
        # Validate the supplied dates.
        if start_date:
            date.fromisoformat(start_date)

        if end_date:
            date.fromisoformat(end_date)

        if start_date and end_date and start_date > end_date:
            flash("Start date cannot be after end date.", "error")
            return redirect(url_for("reports"))

        connection = get_db_connection()
        cursor = connection.cursor(dictionary=True)

        query = """
            SELECT
                sm.id,
                p.name AS product_name,
                p.category,
                sm.movement_type,
                sm.quantity,
                sm.reason,
                sm.created_at
            FROM stock_movements sm
            JOIN products p ON p.id = sm.product_id
            WHERE 1 = 1
        """

        parameters = []

        if start_date:
            query += " AND sm.created_at >= %s"
            parameters.append(start_date)

        if end_date:
            query += " AND sm.created_at < DATE_ADD(%s, INTERVAL 1 DAY)"
            parameters.append(end_date)

        query += " ORDER BY sm.created_at DESC, sm.id DESC"

        cursor.execute(query, tuple(parameters))
        movements = cursor.fetchall()

        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow([
            "Movement ID",
            "Product Name",
            "Category",
            "Movement Type",
            "Quantity",
            "Reason",
            "Date and Time"
        ])

        for movement in movements:
            writer.writerow([
                movement["id"],
                movement["product_name"],
                movement["category"],
                movement["movement_type"],
                movement["quantity"],
                movement["reason"] or "Not specified",
                movement["created_at"]
            ])

        csv_content = "\ufeff" + output.getvalue()
        output.close()

        filename = "stock_movement_report.csv"

        return Response(
            csv_content,
            mimetype="text/csv",
            headers={
                "Content-Disposition": f"attachment; filename={filename}"
            }
        )

    except ValueError:
        flash("Please enter valid dates.", "error")
        return redirect(url_for("reports"))

    except Exception:
        app.logger.exception("Could not export stock movement report")
        flash("Unable to export the stock movement report.", "error")
        return redirect(url_for("reports"))

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None and connection.is_connected():
            connection.close()



if __name__ == "__main__":
    app.run(debug=True)

